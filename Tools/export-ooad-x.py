# -*- coding: utf-8 -*-
"""Export docs/ooad chapter X markdown to docx."""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
CYAN = RGBColor(0x1A, 0x5F, 0x7A)
CODE = RGBColor(0x1A, 0x3A, 0x5C)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "ooad" / "X-Phan-tich-thiet-ke-huong-doi-tuong.md"
OUT = ROOT / "docs" / "ooad" / "X-Phan-tich-thiet-ke-huong-doi-tuong.docx"


def set_run_font(run, size=12, bold=False, italic=False, color=None, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_runs(paragraph, text, size=12, color=None):
    """Render **bold** and `code` inside a paragraph."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, bold=False, color=CODE, name="Consolas")
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size, color=color)


def add_p(doc, text, *, size=12, first_line=True, space_after=8, align=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.15
    pf.first_line_indent = Cm(1) if first_line else Cm(0)
    add_runs(p, text, size=size, color=color)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(16 if level == 1 else 12)
    pf.space_after = Pt(8)
    pf.first_line_indent = Cm(0)
    pf.line_spacing = 1.15
    add_runs(p, text, size=16 if level == 1 else 14, color=NAVY)
    for run in p.runs:
        run.bold = True
        run.font.color.rgb = NAVY
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): hex_color, qn("w:val"): "clear"})
    tcPr.append(shd)


def add_table(doc, headers, rows, caption):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(paragraph, header, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
        for run in paragraph.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(cell, "1B2A4A")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_runs(paragraph, val, size=10)
            if r_i % 2 == 1:
                shade_cell(cell, "EEF3F8")
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.first_line_indent = Cm(0)
    add_runs(cap, caption, size=11, color=CYAN)
    for run in cap.runs:
        run.bold = True
        run.font.color.rgb = CYAN
    return table


def strip_link(text):
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [strip_link(c.strip()) for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) or c.replace("-", "") == "" for c in cells):
            if not re.match(r"^[\s|:-]+$", lines[i].strip()):
                rows.append(cells)
        i += 1
    return rows, i


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    table_index = 0
    captions = [
        "Bảng X.1 — Phạm vi OOAD bám GDD và lớp đang chạy",
        "Bảng X.2 — Danh từ trong GDD và lớp giữ trạng thái",
    ]
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("# "):
            add_heading(doc, strip_link(stripped[2:].strip()), 1)
            i += 1
            continue
        if stripped.startswith("## "):
            add_heading(doc, strip_link(stripped[3:].strip()), 2)
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if len(rows) >= 2:
                caption = captions[table_index] if table_index < len(captions) else "Bảng"
                table_index += 1
                add_table(doc, rows[0], rows[1:], caption)
            continue
        if re.match(r"^\d+\.\s", stripped):
            add_p(doc, strip_link(stripped), first_line=False, space_after=6)
            i += 1
            continue
        if stripped.startswith("- "):
            add_p(doc, strip_link(stripped[2:].strip()), first_line=False, space_after=4)
            i += 1
            continue
        add_p(doc, strip_link(stripped), first_line=not stripped.startswith("**"))
        i += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(OUT)
        print(OUT)
    except PermissionError:
        for suffix in ("-cap-nhat", "-FR"):
            alt = OUT.with_name(OUT.stem + suffix + ".docx")
            try:
                doc.save(alt)
                print(alt)
                return
            except PermissionError:
                continue
        raise


if __name__ == "__main__":
    build()
